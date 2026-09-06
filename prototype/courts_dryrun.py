"""Proof-of-concept: COURTS consignment file -> the 12 invoices it should produce.

Not the app. It exists to prove the plan's arithmetic on a real file before
anyone writes the real thing. No Xero calls; prints what it would post.

    pip install openpyxl
    python prototype/courts_dryrun.py <ConsignmentPO*.xlsx> [master_price_list.xlsx]

If the master price list is given, each SKU is routed to its Xero entity by
the master's BRAND column (Table Matters -> AGPL, everything else -> SGPL) and
one invoice is produced per (entity, store). Without it, SKUs are not routed
and a single entity is assumed. The master is never committed to this repo.
"""
import sys
from collections import defaultdict
from datetime import date
from decimal import Decimal

import openpyxl

from gst import line_exclusive, invoice_totals, rate_for, r2

EXPECTED_COMMISSION_TIERS = (Decimal("0.30"), Decimal("0.35"))
TIER_TOLERANCE = Decimal("0.002")
AGPL_BRANDS = {"TABLE MATTERS"}


def load_brand_map(path):
    """From the master price list, build two lookups:
       brand_by_sku   : our SKU -> BRAND
       sku_by_courts  : COURTS' article code (their 'Item No_') -> our SKU
    Master Price List 4.0 keeps a 'COURTS SKU' column per row - that column
    *is* the customer alias table, maintained by the business already."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    alias = {}
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        # header may sit a few rows down (Master Price List 4.0 uses row 4)
        hdr = None
        for _ in range(10):
            cand = [str(h).strip().upper() if h else "" for h in next(rows, ())]
            if "SKU" in cand and "BRAND" in cand:
                hdr = cand
                break
        if hdr is None:
            continue
        si, bi = hdr.index("SKU"), hdr.index("BRAND")
        ci = hdr.index("COURTS SKU") if "COURTS SKU" in hdr else None
        for r in rows:
            if len(r) > max(si, bi) and r[si] and r[bi]:
                sku = str(r[si]).strip().upper()
                out.setdefault(sku, str(r[bi]).strip().upper())
                if ci is not None and len(r) > ci and r[ci]:
                    alias.setdefault(str(r[ci]).strip().upper(), sku)
    return {"brand_by_sku": out, "sku_by_courts": alias}


def resolve(row, master):
    """Match ladder (plan section 7): tier 2 customer alias, then tier 3 our
    SKU verbatim. Returns (our_sku, entity)."""
    if master is None:
        return str(row["Model"]), "SGPL"
    item_no = str(row["Item No_"]).strip().upper()
    model = str(row["Model"]).strip().upper()
    sku = master["sku_by_courts"].get(item_no)          # tier 2
    if sku is None and model in master["brand_by_sku"]:
        sku = model                                      # tier 3
    if sku is None:
        return model, "UNMAPPED"   # the real app hard-stops here
    brand = master["brand_by_sku"].get(sku)
    return sku, ("AGPL" if brand in AGPL_BRANDS else "SGPL")


def load(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Sheet1"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, [c.value for c in ws[r]]))
            for r in range(2, ws.max_row + 1)]
    # profiles/courts.yaml row_filters
    return [r for r in rows if r["Entry No_"] not in (None, "", 0)
            and r["Invoiced Quantity"]]


def main(path, master=None):
    rows = load(path)
    rate = rate_for(date(2026, 8, 31))
    brand_map = load_brand_map(master) if master else None
    stores = defaultdict(lambda: defaultdict(lambda: [Decimal(0), Decimal(0), None]))
    ar = {}
    warnings = []

    for row in rows:
        sku, entity = resolve(row, brand_map)
        if entity == "UNMAPPED":
            warnings.append(f"{row['Item No_']} / {row['Model']}: not in "
                            f"master (COURTS SKU or SKU) -> cannot route")
        store = (entity, str(row["Location Code"]))
        ar.setdefault(store, set()).add(row["PONUMBER"])
        qty = Decimal(str(row["Invoiced Quantity"]))
        cost = Decimal(str(row["CostPrice"]))
        retail = Decimal(str(row["Unit Price"]))

        # Proof 1: commission is VERIFIED, never calculated (plan section 6.1)
        if retail:
            implied = Decimal(1) - cost / (retail / (Decimal(1) + rate))
            if not any(abs(implied - t) <= TIER_TOLERANCE
                       for t in EXPECTED_COMMISSION_TIERS):
                warnings.append(f"{row['Model']}: implied commission "
                                f"{implied:.4f} is not a known tier")

        bucket = stores[store][sku]
        bucket[0] += qty
        bucket[1] += qty * cost
        bucket[2] = row["Description"]

    grand = defaultdict(lambda: {"net": Decimal(0), "tax": Decimal(0),
                                 "total": Decimal(0), "n": 0})
    for (entity, store) in sorted(stores):
        refs = ar[(entity, store)]
        assert len(refs) == 1, f"store {store} has multiple AR numbers: {refs}"
        ref = refs.pop()
        lines = [line_exclusive(v[1], rate)
                 for v in stores[(entity, store)].values()]
        t = invoice_totals(lines)
        print(f"{entity}  store {store}  {ref}  {len(lines):2} lines   "
              f"net {t['net']:>9}  GST {t['tax']:>7}  total {t['total']:>9}")
        for k in t:
            grand[entity][k] += t[k]
        grand[entity]["n"] += 1

    print()
    for entity, g in sorted(grand.items()):
        print(f"{entity}: {g['n']} invoices   net {g['net']}   "
              f"GST {g['tax']}   total {g['total']}")
    print(f"rows in: {len(rows)}   rows invoiced: "
          f"{sum(len(v) for v in stores.values())} aggregated lines")
    for w in sorted(set(warnings)):
        print(f"  ⚠ {w}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
